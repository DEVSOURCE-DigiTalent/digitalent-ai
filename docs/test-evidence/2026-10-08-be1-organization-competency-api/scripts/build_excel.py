"""Builds the evidence workbook from results.json (E2E), auto.json (automated checks) and rules.json (rule scan)."""
import json
import sys
from datetime import datetime, timezone, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from PIL import Image as PILImage

ROOT = Path(sys.argv[1])          # evidence folder
META = json.loads(sys.argv[2])    # {"branch", "commit", "date"}
IMAGES = ROOT / 'images'
VN = timezone(timedelta(hours=7))

e2e = json.loads((ROOT / 'results.json').read_text(encoding='utf-8'))
auto = json.loads((ROOT / 'auto.json').read_text(encoding='utf-8'))
rules = json.loads((ROOT / 'rules.json').read_text(encoding='utf-8'))
findings = json.loads((ROOT / 'findings.json').read_text(encoding='utf-8'))
cases = e2e + auto

THIN = Side(style='thin', color='B7B7B7')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FILL = PatternFill('solid', fgColor='1F3864')
PASS_FILL = PatternFill('solid', fgColor='C6EFCE')
FAIL_FILL = PatternFill('solid', fgColor='FFC7CE')
LABEL_FILL = PatternFill('solid', fgColor='F2F2F2')
WRAP = Alignment(wrap_text=True, vertical='top')
CENTER = Alignment(horizontal='center', vertical='top', wrap_text=True)


def when(case):
    return datetime.fromisoformat(case['at'].replace('Z', '+00:00')).astimezone(VN).strftime('%d/%m/%Y %H:%M:%S')


def api_text(case):
    calls = case.get('api') or []
    # keep the business calls; drop polling noise
    calls = [c for c in calls if not c['path'].startswith('/api/v1/notifications')]
    seen, lines = set(), []
    for c in calls:
        key = (c['method'], c['path'].split('?')[0], c['status'])
        if key in seen:
            continue
        seen.add(key)
        lines.append(f"{c['method']} {c['path']} → {c['status']}")
    return '\n'.join(lines) if lines else (case.get('apiNote') or '—')


wb = Workbook()
ws = wb.active
ws.title = 'Tong hop'

ws['A1'] = 'BÁO CÁO KIỂM THỬ BE1 – API ↔ FRONTEND: TỔ CHỨC, NĂNG LỰC, NHÂN VIÊN (OW-01…OW-22, AUTH-05, EM-01…EM-18)'
ws['A1'].font = Font(bold=True, size=14, color='1F3864')
ws['A2'] = (f"Branch {META['branch']} @ {META['commit']} · FE http://localhost:5173 (Vite, VITE_USE_MOCK=false) · "
            f"API http://localhost:5000 (Development) · PostgreSQL 17.6, DB riêng digitalent_fe_evidence (migration + seed mới) · Chrome headless 1440×900 · {META['date']}")
ws['A2'].font = Font(italic=True, size=10, color='595959')

last_row = 9 + len(cases) - 1
summary = [
    ('Tổng số ca kiểm thử', f'=COUNTA(B9:B{last_row})'),
    ('Số ca PASS', f'=COUNTIF(L9:L{last_row},"PASS")'),
    ('Số ca FAIL', f'=COUNTIF(L9:L{last_row},"FAIL")'),
    ('Tỷ lệ PASS', '=IF(B4=0,0,B5/B4)'),
]
for i, (label, formula) in enumerate(summary, start=4):
    ws.cell(i, 1, label).font = Font(bold=True)
    c = ws.cell(i, 2, formula)
    c.font = Font(bold=True)
    if label == 'Tỷ lệ PASS':
        c.number_format = '0%'

headers = ['STT', 'Mã TC', 'Màn hình', 'Ca kiểm thử', 'URL / thao tác', 'Vai trò', 'Tiền điều kiện / Dữ liệu',
           'Các bước thực hiện', 'Kết quả mong đợi', 'Kết quả thực tế', 'API backend đã gọi (thực tế)', 'Trạng thái', 'Evidence', 'Thời gian thực hiện']
widths = [5, 8, 18, 30, 26, 15, 28, 40, 40, 36, 42, 10, 11, 18]
for col, (h, w) in enumerate(zip(headers, widths), start=1):
    c = ws.cell(8, col, h)
    c.font = Font(bold=True, color='FFFFFF')
    c.fill = HEADER_FILL
    c.alignment = CENTER
    c.border = BORDER
    ws.column_dimensions[get_column_letter(col)].width = w
ws.freeze_panes = 'C9'

for i, case in enumerate(cases, start=1):
    r = 8 + i
    steps = '\n'.join(f'{n}. {s}' for n, s in enumerate(case.get('steps', []), start=1))
    values = [i, case['id'], case['screen'], case['title'], case['url'], case['role'], case.get('precondition', ''),
              steps, case['expected'], case.get('actual', ''), api_text(case), case['status'], f"Xem {case['id']}", when(case)]
    for col, v in enumerate(values, start=1):
        c = ws.cell(r, col, v)
        c.alignment = CENTER if col in (1, 2, 12, 13) else WRAP
        c.border = BORDER
    status = ws.cell(r, 12)
    status.fill = PASS_FILL if case['status'] == 'PASS' else FAIL_FILL
    status.font = Font(bold=True, color='006100' if case['status'] == 'PASS' else '9C0006')
    link = ws.cell(r, 13)
    link.hyperlink = f"#'{case['id']}'!A1"
    link.font = Font(color='0563C1', underline='single')

notes_row = last_row + 2
notes = [
    'Ghi chú',
    '1. Mỗi ca có một sheet cùng tên mã TC: thông tin ca, chú thích từng vùng khoanh đỏ (số trên khung đỏ ↔ số trong chú thích) và ảnh chụp màn hình.',
    '2. Các ca UI (TC01–TC53) chạy tuần tự trên cùng một DB mới tạo từ migration + seed Development; ca sau dùng dữ liệu của ca trước (VD TC04–TC06 dùng lời mời tạo ở TC02, TC51–TC53 dùng bài nộp ở TC50). Nhóm: TC01–TC25 Tổ chức (Owner), TC26–TC39 Năng lực (Owner), TC40–TC50 + TC53 Nhân viên, TC51–TC52 Quản lý chấm minh chứng.',
    '3. Cột "API backend đã gọi" ghi lại request thật trình duyệt gửi tới backend (method, đường dẫn, HTTP status) trong lúc thực hiện ca — chứng minh dữ liệu không đến từ mock.',
    '4. Mật khẩu/token không hiển thị trong ảnh (ô mật khẩu bị che, token chỉ nằm trong URL kích hoạt). Tài khoản demo theo seed Development; tài khoản được mời chỉ tồn tại trong DB kiểm thử.',
    '5. AUTO01–AUTO03: frontend (Vitest, tsc/lint/build); AUTO04–AUTO05: backend (dotnet build/test trên PostgreSQL). Sheet "Ra soat rule": đối chiếu rule dự án với thay đổi của nhánh. Thư mục scripts/: công cụ chạy lại toàn bộ evidence.',
    '6. Sau khi đăng nhập, script tải lại trang qua origin 127.0.0.1 rồi quay về: Chrome headless ngừng nhận chuột/phím trong tab đăng nhập khi vào cổng nhân viên. Đã kiểm tra trên trình duyệt thật: người dùng không gặp hiện tượng này.',
]
for k, text in enumerate(notes):
    c = ws.cell(notes_row + k, 1, text)
    c.font = Font(bold=(k == 0))

find_row = notes_row + len(notes) + 1
ws.cell(find_row, 1, 'Phát hiện trong quá trình kiểm thử').font = Font(bold=True, color='C00000')
find_headers = ['Mã', 'Mức độ', 'Vấn đề', 'Bằng chứng', 'Ảnh hưởng', 'Đề xuất']
find_cols = [2, 3, 4, 7, 9, 11]   # spread over the wide summary columns
for h, col in zip(find_headers, find_cols):
    c = ws.cell(find_row + 1, col, h)
    c.font = Font(bold=True, color='FFFFFF')
    c.fill = HEADER_FILL
    c.border = BORDER
for k, f in enumerate(findings, start=2):
    for key, col in zip(['id', 'severity', 'title', 'evidence', 'impact', 'action'], find_cols):
        c = ws.cell(find_row + k, col, f[key])
        c.alignment = WRAP
        c.border = BORDER
    ws.row_dimensions[find_row + k].height = 75

# ─────────────── one sheet per case ───────────────
for case in cases:
    s = wb.create_sheet(case['id'])
    s.column_dimensions['A'].width = 24
    s.column_dimensions['B'].width = 150
    s['A1'] = f"{case['id']} – {case['title']}"
    s['A1'].font = Font(bold=True, size=13, color='1F3864')
    back = s['A2']
    back.value = '← Quay lại Tổng hợp'
    back.hyperlink = "#'Tong hop'!A1"
    back.font = Font(color='0563C1', underline='single')
    info = [
        ('Màn hình', case['screen']), ('Vai trò', case['role']), ('URL / thao tác', case['url']),
        ('Tiền điều kiện', case.get('precondition', '')),
        ('Các bước', '\n'.join(f'{n}. {x}' for n, x in enumerate(case.get('steps', []), start=1))),
        ('Kết quả mong đợi', case['expected']), ('Kết quả thực tế', case.get('actual', '')),
        ('API backend đã gọi', api_text(case)), ('Trạng thái', case['status']), ('Thời gian', when(case)),
        ('Chú thích khoanh đỏ', '\n'.join(case.get('notes', [])) or '—'),
    ]
    row = 4
    for label, value in info:
        a = s.cell(row, 1, label)
        a.font = Font(bold=True)
        a.fill = LABEL_FILL
        a.alignment = WRAP
        a.border = BORDER
        b = s.cell(row, 2, value)
        b.alignment = WRAP
        b.border = BORDER
        lines = max(1, str(value).count('\n') + 1, len(str(value)) // 140 + 1)
        s.row_dimensions[row].height = min(400, 15 * lines + 4)
        if label == 'Trạng thái':
            b.fill = PASS_FILL if value == 'PASS' else FAIL_FILL
            b.font = Font(bold=True, color='006100' if value == 'PASS' else '9C0006')
        if label == 'Chú thích khoanh đỏ':
            b.font = Font(color='C00000')
        row += 1

    row += 1
    images = [f'{x}.png' for x in case.get('extraImages', [])] + [f"{case['id']}.png"]
    for name in images:
        path = IMAGES / name
        if not path.exists():
            continue
        s.cell(row, 1, f'Ảnh: {name}').font = Font(bold=True, italic=True)
        row += 1
        with PILImage.open(path) as im:
            w, h = im.size
        scale = min(1.0, 1150 / w)
        img = XLImage(str(path))
        img.width, img.height = int(w * scale), int(h * scale)
        s.add_image(img, f'A{row}')
        row += int(img.height / 20) + 2   # default row height ≈ 20 px

# ─────────────── rule scan sheet ───────────────
rs = wb.create_sheet('Ra soat rule', 1)
rs['A1'] = 'RÀ SOÁT RULE DỰ ÁN – THAY ĐỔI CỦA NHÁNH'
rs['A1'].font = Font(bold=True, size=13, color='1F3864')
rs['A2'] = 'Nguồn rule: CLAUDE.md (Backend, Frontend, Git, Testing), backend/DEVELOPER_GUIDE.md, frontend/DEVELOPER_GUIDE.md. Phạm vi: diff origin/develop…nhánh (backend + frontend + docs).'
rs['A2'].font = Font(italic=True, size=10, color='595959')
rule_headers = ['STT', 'Rule', 'Nguồn', 'Cách kiểm tra', 'Kết quả', 'Đạt']
rule_widths = [5, 48, 26, 52, 60, 8]
for col, (h, w) in enumerate(zip(rule_headers, rule_widths), start=1):
    c = rs.cell(4, col, h)
    c.font = Font(bold=True, color='FFFFFF')
    c.fill = HEADER_FILL
    c.alignment = CENTER
    c.border = BORDER
    rs.column_dimensions[get_column_letter(col)].width = w
for i, rule in enumerate(rules, start=1):
    vals = [i, rule['rule'], rule['source'], rule['check'], rule['result'], rule['ok']]
    for col, v in enumerate(vals, start=1):
        c = rs.cell(4 + i, col, v)
        c.alignment = CENTER if col in (1, 6) else WRAP
        c.border = BORDER
    ok = rs.cell(4 + i, 6)
    ok.fill = PASS_FILL if rule['ok'] == 'ĐẠT' else FAIL_FILL
    ok.font = Font(bold=True)

out = ROOT / META.get('file', 'Evidence_BE1.xlsx')
wb.save(out)
print(out, len(cases), 'cases')

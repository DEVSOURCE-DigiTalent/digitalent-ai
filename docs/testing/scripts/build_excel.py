"""Build the EM-01..EM-18 test evidence workbook from results.json, the screenshots and the automated test reports."""
import json
import os
import sys
import xml.etree.ElementTree as ET
from collections import Counter, OrderedDict
from datetime import datetime

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from PIL import Image

SP, OUT_XLSX, COMMIT, BRANCH = sys.argv[1], sys.argv[2], sys.argv[3], sys.argv[4]
EVIDENCE = os.path.join(SP, 'evidence')
THUMBS = os.path.join(EVIDENCE, 'thumbs')
os.makedirs(THUMBS, exist_ok=True)
results = json.load(open(os.path.join(EVIDENCE, 'results.json'), encoding='utf-8'))
TESTER = 'Nguyễn Thiện Hoàng'
RUN_DATE = datetime.fromisoformat(results['startedAt']).strftime('%d/%m/%Y')

FONT = 'Arial'
BLUE = '1F4E79'
thin = Side(style='thin', color='BFBFBF')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HEADER_FILL = PatternFill('solid', start_color=BLUE)
SECTION_FILL = PatternFill('solid', start_color='DDEBF7')
WRAP_TOP = Alignment(wrap_text=True, vertical='top')
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)


def font(**kw):
    return Font(name=FONT, size=kw.pop('size', 10), **kw)


def header(ws, row, titles, widths=None):
    for col, title in enumerate(titles, start=1):
        cell = ws.cell(row=row, column=col, value=title)
        cell.font = font(bold=True, color='FFFFFF')
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = BORDER
    if widths:
        for col, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[row].height = 30


def body(cell, bold=False, align=WRAP_TOP):
    cell.font = font(bold=bold)
    cell.alignment = align
    cell.border = BORDER


def status_rules(ws, rng):
    ws.conditional_formatting.add(rng, CellIsRule(operator='equal', formula=['"Pass"'],
                                                  fill=PatternFill('solid', start_color='C6EFCE'), font=Font(name=FONT, color='006100', bold=True)))
    ws.conditional_formatting.add(rng, CellIsRule(operator='equal', formula=['"Fail"'],
                                                  fill=PatternFill('solid', start_color='FFC7CE'), font=Font(name=FONT, color='9C0006', bold=True)))


wb = Workbook()
wb.calculation.fullCalcOnLoad = True  # tính lại công thức khi mở bằng Excel

# ─────────────────────────── Sheet: Chi tiết bước & Evidence ───────────────────────────
steps_ws = wb.active
steps_ws.title = 'Chi tiết bước & Evidence'
STEP_HEADERS = ['Mã TC', 'Bước', 'Thao tác', 'Dữ liệu nhập', 'Kết quả mong đợi', 'Kết quả thực tế', 'Trạng thái', 'URL sau bước', 'Tên ảnh', 'Ảnh evidence']
steps_ws['A1'] = 'CHI TIẾT TỪNG BƯỚC VÀ ẢNH EVIDENCE — EM-01 → EM-18'
steps_ws['A1'].font = font(bold=True, size=14, color=BLUE)
steps_ws['A2'] = 'Trong ảnh: khung đỏ = phần tử được thao tác/kiểm tra ở bước đó; nhãn vàng góc phải = mã test case và nội dung bước. Bước "bấm" được chụp ngay trước khi bấm.'
steps_ws['A2'].font = font(italic=True, color='595959')
header(steps_ws, 4, STEP_HEADERS, [13, 6, 30, 26, 32, 40, 11, 30, 22, 100])
IMG_W, IMG_H = 720, 450
row = 5
step_rows = []
for case in results['cases']:
    for step in case['steps']:
        values = [case['id'], step['no'], step['action'], step['data'], step['expected'], step['actual'], step['status'], step['url'], step['shot']]
        for col, value in enumerate(values, start=1):
            body(steps_ws.cell(row=row, column=col, value=value), align=CENTER if col in (1, 2, 7) else WRAP_TOP)
        body(steps_ws.cell(row=row, column=10))
        source = os.path.join(EVIDENCE, 'shots', step['shot'])
        if os.path.exists(source):
            thumb = os.path.join(THUMBS, step['shot'].replace('.png', '.jpg'))
            with Image.open(source) as image:
                image.convert('RGB').resize((IMG_W, IMG_H), Image.LANCZOS).save(thumb, 'JPEG', quality=82, optimize=True)
            picture = XLImage(thumb)
            picture.width, picture.height = IMG_W, IMG_H
            steps_ws.add_image(picture, f'J{row}')
        steps_ws.row_dimensions[row].height = IMG_H * 0.75 + 8
        step_rows.append(row)
        row += 1
last_step_row = row - 1
steps_ws.freeze_panes = 'C5'
status_rules(steps_ws, f'G5:G{last_step_row}')
dv = DataValidation(type='list', formula1='"Pass,Fail,Blocked"', allow_blank=False)
steps_ws.add_data_validation(dv)
dv.add(f'G5:G{last_step_row}')
STEP_ID = f"'Chi tiết bước & Evidence'!$A$5:$A${last_step_row}"
STEP_STATUS = f"'Chi tiết bước & Evidence'!$G$5:$G${last_step_row}"

# ─────────────────────────── Sheet: Test cases ───────────────────────────
cases_ws = wb.create_sheet('Test cases', 0)
cases_ws['A1'] = 'DANH SÁCH TEST CASE — MÀN HÌNH CÁ NHÂN NHÂN VIÊN (EM-01 → EM-18)'
cases_ws['A1'].font = font(bold=True, size=14, color=BLUE)
cases_ws['A2'] = f'Trạng thái và số bước được tính bằng công thức từ sheet "Chi tiết bước & Evidence". Ngày test: {RUN_DATE} · Người test: {TESTER}'
cases_ws['A2'].font = font(italic=True, color='595959')
CASE_HEADERS = ['STT', 'Mã TC', 'Màn hình', 'Tên test case', 'Mục tiêu', 'Tiền điều kiện', 'Dữ liệu test', 'Số bước',
                'Số bước Fail', 'Kết quả mong đợi', 'Trạng thái', 'Ngày test', 'Người test']
header(cases_ws, 4, CASE_HEADERS, [6, 13, 24, 34, 36, 32, 30, 8, 9, 36, 11, 11, 18])
case_first = 5
for index, case in enumerate(results['cases'], start=1):
    r = case_first + index - 1
    values = [index, case['id'], case['screen'], case['name'], case['goal'], case['precondition'], case['data'] or '—']
    for col, value in enumerate(values, start=1):
        body(cases_ws.cell(row=r, column=col, value=value), align=CENTER if col in (1, 2) else WRAP_TOP)
    body(cases_ws.cell(row=r, column=8, value=f'=COUNTIF({STEP_ID},B{r})'), align=CENTER)
    body(cases_ws.cell(row=r, column=9, value=f'=COUNTIFS({STEP_ID},B{r},{STEP_STATUS},"<>Pass")'), align=CENTER)
    body(cases_ws.cell(row=r, column=10, value=case['expected']))
    body(cases_ws.cell(row=r, column=11, value=f'=IF(H{r}=0,"Blocked",IF(I{r}=0,"Pass","Fail"))'), bold=True, align=CENTER)
    body(cases_ws.cell(row=r, column=12, value=RUN_DATE), align=CENTER)
    body(cases_ws.cell(row=r, column=13, value=TESTER), align=CENTER)
    cases_ws.row_dimensions[r].height = 62
case_last = case_first + len(results['cases']) - 1
cases_ws.freeze_panes = 'C5'
status_rules(cases_ws, f'K5:K{case_last}')
cases_ws.auto_filter.ref = f'A4:M{case_last}'

# ─────────────────────────── Sheet: Test tự động ───────────────────────────
auto_ws = wb.create_sheet('Test tự động', 1)
auto_ws['A1'] = 'TEST TỰ ĐỘNG CHO PHẦN EM-01 → EM-18'
auto_ws['A1'].font = font(bold=True, size=14, color=BLUE)
auto_ws['A2'] = 'Nguồn số liệu: báo cáo chạy thật trên nhánh ' + BRANCH + ' — backend: dotnet test (xUnit, PostgreSQL 16 + EF InMemory); frontend: vitest (Testing Library).'
auto_ws['A2'].font = font(italic=True, color='595959')
header(auto_ws, 4, ['STT', 'Tầng', 'File test', 'Màn hình', 'Nội dung kiểm thử', 'Số test', 'Đạt', 'Không đạt', 'Commit', 'Kết quả'],
       [6, 11, 44, 13, 60, 8, 8, 10, 11, 11])

trx = ET.parse(os.path.join(SP, 'test-results', 'be-tests.trx'))
ns = {'t': trx.getroot().tag.split('}')[0].strip('{')}
backend = Counter()
backend_failed = Counter()
for result in trx.getroot().iter(f"{{{ns['t']}}}UnitTestResult"):
    cls = result.get('testName').split('(')[0].rsplit('.', 1)[0]  # bỏ tham số Theory (có thể chứa dấu chấm) rồi bỏ tên method
    backend[cls] += 1
    if result.get('outcome') != 'Passed':
        backend_failed[cls] += 1
vitest = json.load(open(os.path.join(SP, 'vitest_tests_branch.json'), encoding='utf-8'))
frontend = {}
for file_result in vitest['testResults']:
    name = file_result['name'].replace(chr(92), '/').split('/frontend/')[-1]
    frontend[name] = (len(file_result['assertionResults']), sum(1 for a in file_result['assertionResults'] if a['status'] != 'passed'))

AUTO = [
    ('Backend', 'DigiTalent.Tests.Me.MyCompetencyOverviewTests', 'EM-01 → 03', 'Dashboard, hồ sơ năng lực, khoảng trống trên dữ liệu demo thật; không có vị trí → lý do bỏ qua; không có hồ sơ → 403', '3523e0b'),
    ('Backend', 'DigiTalent.Tests.Me.MyGrowthJourneyTests', 'EM-04, 05, 18', 'Dòng thời gian minh chứng chỉ của mình; lộ trình (khóa tiên quyết đứng trước); thành tựu và chứng chỉ', '4396b84'),
    ('Backend', 'DigiTalent.Tests.Me.MyCourseLearningTests', 'EM-06 → 08', 'Khóa học của tôi, chi tiết khóa, mở/hoàn thành bài, tiến độ, mở khóa bài cuối khóa', '50bbf21'),
    ('Backend', 'DigiTalent.Tests.Me.MyAttemptSessionTests', 'EM-10 → 13', 'Tự lưu đáp án, khôi phục khi tải lại, chặn lưu sau khi nộp, lịch sử lọc/phân trang', 'def57d6'),
    ('Backend', 'DigiTalent.Tests.Me.MyTaskEvidenceTests', 'EM-14 → 17', 'Chi tiết nhiệm vụ, tải tệp vào thư mục nhiệm vụ, nộp + tải lại tệp, chặn nhiệm vụ/tệp của người khác', '9357b94'),
    ('Backend', 'DigiTalent.Tests.Me.MyTaskInputValidationTests', 'EM-16', 'Kiểm tra dữ liệu: mô tả ≥ 20 ký tự, link http(s), tối đa 10 link/tệp, định dạng và dung lượng tệp', '9357b94'),
    ('Backend', 'DigiTalent.Tests.Me.MyAssessmentFlowTests', 'EM-09 → 13', 'Khóa bài cuối khóa, chấm điểm, hết giờ, giới hạn lượt, cấp chứng chỉ (từ PR #53)', '0e80637'),
    ('Backend', 'DigiTalent.Tests.Me.MyTaskAndCourseTests', 'EM-06 → 17', 'Nhiệm vụ của mình, nộp/nộp lại, tự ghi danh khóa tiên quyết (từ PR #53)', '0e80637'),
    ('Backend', 'DigiTalent.Tests.Me.MyHelpersTests', 'EM-05, 16', 'Đọc rubric, liên kết minh chứng, sắp xếp lộ trình (từ PR #53)', '0e80637'),
    ('Backend', 'DigiTalent.Tests.Persistence.EmployeeJourneySeederTests', 'Dữ liệu demo', 'Seed dữ liệu demo EM đúng kịch bản, chạy lại không nhân đôi (từ PR #53)', '0e80637'),
    ('Frontend', 'src/features/employee/__tests__/MyDevelopmentScreens.test.tsx', 'EM-01 → 05', 'Đang tải, lỗi + thử lại, lọc minh chứng, ghi danh từ lộ trình', 'bcfc168'),
    ('Frontend', 'src/features/learning/__tests__/EmployeeLearningScreens.test.tsx', 'EM-06 → 08', 'Lọc/tìm khóa học, chi tiết khóa, bài PASS_CHECK, video YouTube, chuyển sang bài đánh giá', 'ca7ffb0'),
    ('Frontend', 'src/features/learning/__tests__/EmployeeAssessmentScreens.test.tsx', 'EM-09 → 13', 'Lọc bài đánh giá, bài bị khóa, tự lưu, hết giờ tự nộp, kết quả ẩn đáp án, lọc lịch sử', '58bcf00'),
    ('Frontend', 'src/features/employee/__tests__/MyTaskScreens.test.tsx', 'EM-14 → 18', 'Lọc nhiệm vụ, tải tệp đúng/sai định dạng, link sai, chờ chấm, chứng chỉ bị thu hồi', '3d26e15'),
    ('Frontend', 'src/features/employee/__tests__/EmployeePages.test.tsx', 'EM-01 → 18', 'Hiển thị dữ liệu thật của 18 màn (từ PR #53)', '672cc81'),
    ('Frontend', 'src/features/employee/__tests__/MyLearningPathPage.test.tsx', 'EM-05', 'Liên kết khóa theo GUID, trạng thái rỗng', '672cc81'),
    ('Frontend', 'src/services/mock/server/__tests__/me-api.test.ts', 'EM-01 → 18', 'Mock server /me/* (chế độ VITE_USE_MOCK)', '672cc81'),
]
auto_first = 5
for index, (layer, name, screens, scope, commit) in enumerate(AUTO, start=1):
    r = auto_first + index - 1
    total, failed = (backend[name], backend_failed[name]) if layer == 'Backend' else frontend.get(name, (0, 0))
    values = [index, layer, name, screens, scope, total, f'=F{r}-H{r}', failed, commit]
    for col, value in enumerate(values, start=1):
        body(auto_ws.cell(row=r, column=col, value=value), align=CENTER if col in (1, 2, 4, 6, 7, 8, 9) else WRAP_TOP)
    body(auto_ws.cell(row=r, column=10, value=f'=IF(F{r}=0,"Blocked",IF(H{r}=0,"Pass","Fail"))'), bold=True, align=CENTER)
    auto_ws.row_dimensions[r].height = 42
auto_last = auto_first + len(AUTO) - 1
total_row = auto_last + 1
auto_ws.cell(row=total_row, column=5, value='Tổng').font = font(bold=True)
for col in (6, 7, 8):
    letter = get_column_letter(col)
    cell = auto_ws.cell(row=total_row, column=col, value=f'=SUM({letter}{auto_first}:{letter}{auto_last})')
    body(cell, bold=True, align=CENTER)
status_rules(auto_ws, f'J5:J{auto_last}')
note_row = total_row + 2
auto_ws.cell(row=note_row, column=1, value='Lệnh chạy lại:').font = font(bold=True)
auto_ws.cell(row=note_row + 1, column=1, value='Backend: cd backend → đặt DIGITALENT_TEST_POSTGRES_CONNECTION (DB tên kết thúc _test) → dotnet test --filter "FullyQualifiedName~DigiTalent.Tests.Me"').font = font()
auto_ws.cell(row=note_row + 2, column=1, value='Frontend: cd frontend → npx vitest run src/features/employee src/features/learning/__tests__/Employee* src/services/mock/server/__tests__/me-api.test.ts').font = font()
auto_ws.freeze_panes = 'A5'

# ─────────────────────────── Sheet: Tổng quan ───────────────────────────
ov = wb.create_sheet('Tổng quan', 0)
ov.column_dimensions['A'].width = 34
ov.column_dimensions['B'].width = 70
ov['A1'] = 'BÁO CÁO KIỂM THỬ — MÀN HÌNH CÁ NHÂN NHÂN VIÊN EM-01 → EM-18'
ov['A1'].font = font(bold=True, size=16, color=BLUE)
ov['A2'] = 'DigiTalent AI · Phần việc BE1 + FE1 · API /api/v1/me/* và 18 trang /enterprise/me/*'
ov['A2'].font = font(italic=True, color='595959')


def section(r, title):
    ov.cell(row=r, column=1, value=title).font = font(bold=True, color='FFFFFF')
    ov.cell(row=r, column=1).fill = HEADER_FILL
    ov.cell(row=r, column=2).fill = HEADER_FILL
    ov.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)


def kv(r, key, value, bold_value=False):
    body(ov.cell(row=r, column=1, value=key), bold=True)
    ov.cell(row=r, column=1).fill = SECTION_FILL
    body(ov.cell(row=r, column=2, value=value), bold=bold_value)


section(4, 'Thông tin kiểm thử')
info = OrderedDict([
    ('Người thực hiện', TESTER),
    ('Ngày thực hiện', RUN_DATE),
    ('Nhánh / commit', f'{BRANCH} @ {COMMIT}'),
    ('Frontend', 'http://localhost:5173 — React 19 + Vite (npm run dev)'),
    ('Backend', 'http://localhost:5000 — ASP.NET Core 8, môi trường Development'),
    ('Cơ sở dữ liệu', 'PostgreSQL 16 (Docker, DB demo local) — dữ liệu demo do DbSeeder tạo lại trước khi chạy'),
    ('Trình duyệt', 'Chromium (Playwright 1.58), cửa sổ 1440×900, tiếng Việt'),
    ('Tài khoản test', 'employee@digitalent.ai (Nhân viên Kế toán) · manager@digitalent.ai (Quản lý) — mật khẩu Admin@1234'),
    ('Cách chạy', 'Test thủ công được thực hiện tự động bằng trình duyệt theo đúng các bước trong sheet "Chi tiết bước & Evidence"; mỗi bước chụp 1 ảnh'),
])
r = 5
for key, value in info.items():
    kv(r, key, value)
    r += 1

r += 1
section(r, 'Kết quả test thủ công (tính bằng công thức)')
r += 1
CASE_STATUS = f"'Test cases'!$K${case_first}:$K${case_last}"
summary = OrderedDict([
    ('Tổng số test case', f"=COUNTA('Test cases'!$B${case_first}:$B${case_last})"),
    ('Test case Pass', f'=COUNTIF({CASE_STATUS},"Pass")'),
    ('Test case Fail', f'=COUNTIF({CASE_STATUS},"Fail")'),
    ('Tổng số bước', f'=COUNTA({STEP_ID})'),
    ('Bước Pass', f'=COUNTIF({STEP_STATUS},"Pass")'),
    ('Bước Fail', f'=COUNTIF({STEP_STATUS},"Fail")'),
])
summary_rows = {}
for key, formula in summary.items():
    kv(r, key, formula, bold_value=True)
    ov.cell(row=r, column=2).alignment = Alignment(horizontal='left', vertical='top')
    summary_rows[key] = r
    r += 1
kv(r, 'Tỉ lệ test case đạt', f"=IF(B{summary_rows['Tổng số test case']}=0,0,B{summary_rows['Test case Pass']}/B{summary_rows['Tổng số test case']})", bold_value=True)
ov.cell(row=r, column=2).number_format = '0.0%'
ov.cell(row=r, column=2).alignment = Alignment(horizontal='left', vertical='top')
r += 2

section(r, 'Kết quả test tự động (từ sheet "Test tự động")')
r += 1
kv(r, 'Tổng số test tự động cho EM', f"='Test tự động'!F{total_row}", bold_value=True)
ov.cell(row=r, column=2).alignment = Alignment(horizontal='left')
r += 1
kv(r, 'Đạt', f"='Test tự động'!G{total_row}", bold_value=True)
ov.cell(row=r, column=2).alignment = Alignment(horizontal='left')
r += 1
kv(r, 'Không đạt', f"='Test tự động'!H{total_row}", bold_value=True)
ov.cell(row=r, column=2).alignment = Alignment(horizontal='left')
r += 2

section(r, 'Cách đọc ảnh evidence')
r += 1
for key, value in [
    ('Khung đỏ', 'Phần tử được thao tác (nút, ô nhập) hoặc được kiểm tra ở bước đó'),
    ('Nhãn vàng (góc phải dưới)', 'Mã test case và nội dung bước, ví dụ "TC-EM16-01 · Bước 3: Chọn tệp sai định dạng virus.exe"'),
    ('Bước "Bấm …"', 'Ảnh chụp ngay TRƯỚC khi bấm để thấy rõ nút; kết quả của cú bấm nằm ở cột "Kết quả thực tế" và ảnh bước kế tiếp'),
    ('Ảnh gốc', 'Ảnh nhúng được thu nhỏ còn 720×450; ảnh gốc 1440×900 tạo lại bằng docs/testing/scripts/em_evidence.py (xem README cùng thư mục)'),
]:
    kv(r, key, value)
    r += 1

for ws in wb.worksheets:
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0

os.makedirs(os.path.dirname(OUT_XLSX), exist_ok=True)
wb.save(OUT_XLSX)

# Giá trị mong đợi của các công thức (để đối chiếu vì máy không có LibreOffice để tính lại)
cases = results['cases']
fails = [c['id'] for c in cases if any(s['status'] != 'Pass' for s in c['steps'])]
steps = [s for c in cases for s in c['steps']]
print(json.dumps({
    'cases': len(cases), 'casesPass': len(cases) - len(fails), 'casesFail': len(fails), 'failCases': fails,
    'steps': len(steps), 'stepsPass': sum(s['status'] == 'Pass' for s in steps),
    'auto': {name: (backend[name] if layer == 'Backend' else frontend.get(name, (0, 0))[0]) for layer, name, *_ in AUTO},
    'xlsxBytes': os.path.getsize(OUT_XLSX),
}, ensure_ascii=False, indent=1))
